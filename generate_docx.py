#!/usr/bin/env python3
"""Generate a DOCX work-item summary from a JSON topic list.

Expects JSON on stdin with this structure:
[
  {
    "nr": 1,
    "datum": "20.02.2026",
    "umfang": "8 Nachr. / ~45k",
    "summary": "Short summary here",
    "note": "Optional italic explanatory note"
  },
  ...
]

Usage:
  python3 generate_docx.py <project_name> <output_path> <input_json_file>
  echo '<json>' | python3 generate_docx.py <project_name> <output_path>
"""

import json
import sys
from copy import deepcopy

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn, nsmap
    from lxml import etree
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn, nsmap
    from lxml import etree


def add_list_table_1_light_style(doc):
    """Inject 'List Table 1 Light' table style via XML (not in python-docx default template)."""
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    style_xml = f"""
    <w:style w:type="table" w:styleId="ListTable1Light"
             xmlns:w="{W}">
      <w:name w:val="List Table 1 Light"/>
      <w:basedOn w:val="TableNormal"/>
      <w:uiPriority w:val="46"/>
      <w:tblPr>
        <w:tblStyleRowBandSize w:val="1"/>
        <w:tblStyleColBandSize w:val="1"/>
      </w:tblPr>
      <w:tblStylePr w:type="firstRow">
        <w:rPr><w:b/></w:rPr>
        <w:tblPr/>
        <w:tcPr>
          <w:tcBorders>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="666666"/>
          </w:tcBorders>
        </w:tcPr>
      </w:tblStylePr>
      <w:tblStylePr w:type="lastRow">
        <w:rPr><w:b/></w:rPr>
        <w:tblPr/>
        <w:tcPr>
          <w:tcBorders>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="666666"/>
          </w:tcBorders>
        </w:tcPr>
      </w:tblStylePr>
      <w:tblStylePr w:type="band1Horz">
        <w:tblPr/>
        <w:tcPr>
          <w:shd w:val="clear" w:color="auto" w:fill="FFFFFF"/>
        </w:tcPr>
      </w:tblStylePr>
    </w:style>
    """
    style_el = etree.fromstring(style_xml)
    doc.styles.element.append(style_el)


def main():
    if len(sys.argv) < 3:
        print("Usage: generate_docx.py <project_name> <output_path> [input_json_file]", file=sys.stderr)
        sys.exit(1)

    project_name = sys.argv[1]
    output_path = sys.argv[2]
    input_path = sys.argv[3] if len(sys.argv) > 3 else None

    if input_path:
        with open(input_path, "r") as f:
            sessions = json.load(f)
    else:
        sessions = json.loads(sys.stdin.read())

    doc = Document()

    # Set Garamond as default font
    style = doc.styles["Normal"]
    style.font.name = "Garamond"
    style.font.size = Pt(11)
    for heading_level in range(1, 4):
        hs = doc.styles[f"Heading {heading_level}"]
        hs.font.name = "Garamond"

    doc.add_heading(f"Arbeitsverlauf \u2013 LLM-Coding \u00d7 {project_name}", level=1)

    add_list_table_1_light_style(doc)
    table = doc.add_table(rows=1, cols=5, style="List Table 1 Light")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Nr.", "Datum", "Modell", "Umfang\n(Nachr. / Tokens)", "Zusammenfassung"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for run in cell.paragraphs[0].runs:
            run.bold = True

    col_widths = [Inches(0.35), Inches(1.0), Inches(0.7), Inches(1.1), Inches(3.35)]

    for row in table.rows:
        for idx, w in enumerate(col_widths):
            row.cells[idx].width = w

    for s in sessions:
        row = table.add_row()
        row.cells[0].text = str(s["nr"])
        row.cells[1].text = s["datum"]
        row.cells[2].text = s.get("modell", "")
        row.cells[3].text = s["umfang"]

        cell = row.cells[4]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(s["summary"])
        run.font.size = Pt(9)

        if s.get("note"):
            p2 = cell.add_paragraph()
            run2 = p2.add_run(s["note"])
            run2.font.size = Pt(8)
            run2.italic = True
            run2.font.color.rgb = RGBColor(128, 128, 128)
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(0)

        for idx, c in enumerate(row.cells):
            c.width = col_widths[idx]
            for para in c.paragraphs:
                para.paragraph_format.space_before = Pt(0)
                para.paragraph_format.space_after = Pt(0)

    # Date range from sessions
    dates = [s["datum"] for s in sessions]
    date_range = f"{dates[0]} \u2013 {dates[-1]}" if len(dates) > 1 else dates[0]

    doc.add_paragraph("")
    footer = doc.add_paragraph(f"{len(sessions)} Einträge | {date_range} | Projekt: {project_name}")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(128, 128, 128)

    doc.save(output_path)
    print(f"Saved to {output_path}")


if __name__ == "__main__":
    main()
